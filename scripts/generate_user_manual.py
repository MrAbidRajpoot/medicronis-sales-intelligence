#!/usr/bin/env python3
"""Generate Medicronis Sales Intelligence User Manual as a Word document."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "Medicronis-User-Manual.docx"
MANUAL_VERSION = "1.3"


@dataclass
class TocEntry:
    level: int
    text: str
    bookmark: str


@dataclass
class ManualBuilder:
    doc: Document
    toc: list[TocEntry] = field(default_factory=list)
    _bookmark_seq: int = 0

    def __post_init__(self) -> None:
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

    def heading(self, text: str, level: int = 1) -> None:
        self._bookmark_seq += 1
        bookmark = f"toc_{self._bookmark_seq}"
        self.toc.append(TocEntry(level=level, text=text, bookmark=bookmark))

        heading = self.doc.add_heading(text, level=level)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bookmark_seq))
        start.set(qn("w:name"), bookmark)
        heading._p.insert(0, start)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bookmark_seq))
        heading._p.append(end)

    def paragraph(self, text: str) -> None:
        self.doc.add_paragraph(text)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def steps(self, items: list[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Number")

    def code_line(self, text: str) -> None:
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(10)

    def table(self, headers: list[str], rows: list[list[str]]) -> None:
        doc = self.doc
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, header in enumerate(headers):
            hdr[i].text = header
            set_cell_shading(hdr[i], "1F4E79")
            for p in hdr[i].paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(255, 255, 255)
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                cell = table.rows[r_idx + 1].cells[c_idx]
                cell.text = val
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(10)
        doc.add_paragraph()

    def page_break(self) -> None:
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)


def set_cell_shading(cell, color_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_hyperlink(paragraph, text: str, bookmark: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def build_toc_elements(doc: Document, toc: list[TocEntry]) -> list:
    """Build TOC section as detached XML elements (hyperlinks, no Word field required)."""
    elements: list = []

    title = doc.add_heading("Table of Contents", level=1)
    elements.append(title._element)

    for entry in toc:
        p = doc.add_paragraph()
        indent = max(0, entry.level - 1) * 0.35
        p.paragraph_format.left_indent = Inches(indent)
        add_hyperlink(p, entry.text, entry.bookmark)
        elements.append(p._element)

    intro = doc.add_paragraph(
        "Click any section name above to jump to that topic. "
        "Section numbers match the headings throughout this manual."
    )
    intro.runs[0].italic = True
    intro.runs[0].font.size = Pt(9)
    intro.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    elements.append(intro._element)

    pb = doc.add_paragraph()
    pb.add_run().add_break(WD_BREAK.PAGE)
    elements.append(pb._element)

    body = doc.element.body
    for el in elements:
        body.remove(el)
    return elements


def add_title_page(doc: Document):
    """Add title page and return the page-break element (TOC inserts after it)."""
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Medicronis Sales Intelligence")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("User Manual")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(68, 114, 196)

    doc.add_paragraph()
    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = desc.add_run(
        "Pharma Distributor Secondary Sales Reporting\n"
        "PDF / Excel Upload → Product Matching → Review → Medicronis SSR Excel (DATA sheet)"
    )
    run.font.size = Pt(12)

    doc.add_paragraph()
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"Version {MANUAL_VERSION}  |  {date.today().strftime('%B %d, %Y')}"
    )
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 100, 100)

    pb = doc.add_paragraph()
    pb.add_run().add_break(WD_BREAK.PAGE)
    return pb._element


def write_content(m: ManualBuilder) -> None:
    # --- 1. Introduction ---
    m.heading("1. Introduction", 1)
    m.paragraph(
        "Medicronis Sales Intelligence is a sales reporting automation platform designed for "
        "pharmaceutical distributors and sales operations teams. The application converts "
        "distributor PDF and Excel sales reports into structured data and generates Secondary "
        "Sales Report (SSR) Excel exports."
    )
    m.heading("1.1 Purpose", 2)
    m.paragraph(
        "The system eliminates manual re-keying of distributor sales reports into spreadsheets. "
        "It supports automated PDF and Excel extraction, intelligent product catalog matching, "
        "human review of exceptions, and consolidated SSR report generation."
    )
    m.heading("1.2 Target Users", 2)
    m.bullets(
        [
            "Medicronis sales operations teams managing distributor networks",
            "Sales managers reviewing distributor performance and sales data",
            "Data entry staff uploading and approving distributor reports",
            "Administrators configuring distributors, products, PDF templates, and Excel column maps",
        ]
    )
    m.heading("1.3 Key Capabilities", 2)
    m.bullets(
        [
            "Upload PDF, Excel (.xlsx), or ZIP archives (mixed PDF + Excel inside ZIP)",
            "Auto-detect distributor from PDF header or Excel filename (with optional manual override)",
            "Excel-only distributors for fragile PDF layouts — PDF uploads blocked; Excel map required",
            "Automatic extraction from 10 PDF layout families (A–J), including Family J line-parser mode",
            "Per-distributor Excel column maps (product name, sales qty, unit price, closing stock, etc.)",
            "Multiple PDF extraction strategies: table parsing, line parser fallback, alternate pdfplumber settings",
            "Fuzzy matching against product catalog with distributor-specific mappings",
            "Product groups managed in the UI and required on every catalog product for SSR grouping",
            "Review queue for unmatched or ambiguous product rows (including create new product)",
            "Approve data and promote to daily sales facts",
            "Generate Medicronis-format SSR Excel (DATA sheet) by as-of date for Day, Week, or Month views",
            "Manage distributors, products, product groups, PDF templates, and Excel maps",
            "Bulk import distributors and products from Excel (.xlsx) templates",
        ]
    )
    m.heading("1.4 End-to-End Workflow Overview", 2)
    m.paragraph("The typical workflow follows these stages:")
    m.steps(
        [
            "Log in to the application (demo password)",
            "Upload a distributor PDF or Excel sales report (distributor auto-detected)",
            "System extracts product rows and matches them to the catalog",
            "Resolve any unmatched rows in the Review Queue (if required)",
            "Approve the document and promote data to sales facts",
            "Generate an SSR report (Day / Week / Month) and download the Excel file",
        ]
    )

    # --- 2. Getting Started ---
    m.heading("2. Getting Started", 1)
    m.heading("2.1 System Requirements", 2)
    m.bullets(
        [
            "Web browser (Chrome, Edge, Firefox, or Safari — latest version recommended)",
            "Network access to the Medicronis application URL",
            "Valid login password (provided by your administrator)",
        ]
    )
    m.paragraph(
        "Note: System administrators also need Node.js, PostgreSQL, and Python 3.11+ "
        "for local development. End users only need a browser."
    )
    m.heading("2.2 Accessing the Application", 2)
    m.paragraph(
        "Open your web browser and navigate to the application URL provided by your administrator."
    )
    m.table(
        ["Environment", "URL"],
        [
            ["Local development", "http://localhost:3000"],
            ["Production / trial demo", "URL provided by your administrator"],
        ],
    )
    m.paragraph(
        "When you first visit the application, you will be redirected to the login page."
    )

    # --- 3. Logging In ---
    m.heading("3. Logging In", 1)
    m.heading("3.1 Login Steps", 2)
    m.steps(
        [
            "Navigate to the application URL in your browser",
            "You will be redirected to the Login page (/login) — titled \"Sign in to demo\"",
            'Enter the shared password in the "Demo Password" field',
            'Click "Access Demo"',
            "Upon successful login, you will be taken to the Dashboard",
        ]
    )
    m.paragraph(
        'If the password is wrong, the page shows: "Invalid password. Use the demo password provided."'
    )
    m.heading("3.2 Login Credentials", 2)
    m.table(
        ["Field", "Value"],
        [
            ["Username", "Not required (single shared password authentication)"],
            ["Demo Password", "Provided by administrator (default: demo)"],
            ["Session duration", "7 days (HTTP-only secure cookie: medicronis-auth)"],
        ],
    )
    m.heading("3.3 Logging Out", 2)
    m.paragraph(
        'To log out, click "Sign out" at the bottom of the sidebar. '
        "This clears your session cookie and returns you to the login page."
    )

    # --- 4. Dashboard ---
    m.heading("4. Dashboard", 1)
    m.paragraph(
        "The Dashboard (/dashboard) is your home screen after login. It provides an "
        "at-a-glance view of sales data and processing activity."
    )
    m.heading("4.1 Key Performance Indicators", 2)
    m.table(
        ["KPI", "Description"],
        [
            ["MTD Total Sales", "Month-to-date approved sales value"],
            ["Pending Documents", "Documents awaiting action (uploaded, processing, extracted, or review required)"],
            ["Product Match Rate", "Percentage of extracted rows matched to the product catalog"],
            ["Active Distributors", "Count of active distributors in the network"],
        ],
    )
    m.heading("4.2 Dashboard Sections", 2)
    m.table(
        ["Section", "Description"],
        [
            ["Sales by Distributor", "Charts for Today, This Week, and This Month"],
            ["Processing Summary", "Counts by status: Approved, Review Required, Processing, Template Mismatch, Failed"],
            ["Template Coverage", "Upload readiness across distributors (PDF template and/or Excel map); links to wizards"],
            ["Recent Activity", "Latest document status changes and SSR generation events"],
        ],
    )
    m.heading("4.3 Using the Dashboard", 2)
    m.bullets(
        [
            'Primary CTA: "Upload PDF" opens the Upload page (PDF, Excel, and ZIP are all supported there)',
            "Review KPIs to monitor overall sales performance",
            "Check Processing Summary and Pending Documents for items awaiting action",
            "Use Template Coverage to find distributors still needing PDF or Excel configuration",
            "If sales charts appear empty, approve at least one document first",
        ]
    )

    # --- 5. Uploading Documents ---
    m.heading("5. Uploading Documents", 1)
    m.paragraph(
        "The Upload page (/upload) is where you submit distributor sales reports "
        "(PDF, Excel, or ZIP) for extraction and SSR generation."
    )
    m.heading("5.1 Before You Upload", 2)
    m.bullets(
        [
            "Have the sales report ready: .pdf, .xlsx, multiple files, or a .zip archive",
            "Set the correct report date (Business date this file represents — cannot be in the future)",
            "For PDF uploads: the matched distributor must have a configured PDF template (unless Excel-only)",
            "For Excel uploads: the matched distributor must have a configured Excel column map",
            "For auto-detect: PDF headers should include the distributor name; Excel filenames should include code or name (e.g. AYAN-TAUNSA.xlsx)",
            "Confirm the PDF worker service is running when uploading PDFs (administrator responsibility)",
        ]
    )
    m.heading("5.2 Supported Upload Formats", 2)
    m.table(
        ["Format", "Description"],
        [
            ["Single PDF", "One distributor sales report; distributor auto-detected from PDF header"],
            ["Single Excel (.xlsx)", "Matched by filename (distributor code or name); requires Excel column map"],
            ["Multiple files", "Select or drag several PDF and/or Excel files; each is processed individually"],
            ["ZIP archive", "One .zip containing PDFs and/or Excel files; each entry is extracted and processed"],
        ],
    )
    m.heading("5.3 Excel-Only Distributors", 2)
    m.paragraph(
        "By default distributors are Both (PDF + Excel; Excel optional). Excel only is set "
        "manually when an operator intentionally wants to block PDF. For Excel-only distributors:"
    )
    m.bullets(
        [
            "Loose PDF uploads are blocked",
            "PDFs inside a ZIP fail per file with an Excel-only error",
            "Upload Excel (.xlsx) or a ZIP of Excel files instead",
            "An Excel column map must be configured before upload (see Section 11)",
            "The Upload distributor dropdown labels these as \"[Excel only]\" / \"— Excel only\"",
        ]
    )
    m.paragraph(
        "See docs/excel-only-distributors.md for historical notes; seed does not force Excel-only."
    )
    m.heading("5.4 Distributor Selection", 2)
    m.paragraph(
        "By default, the system identifies the distributor automatically. Manual selection is optional."
    )
    m.heading("5.4.1 Single PDF Upload", 3)
    m.bullets(
        [
            'Leave the distributor dropdown on "Auto-detect from PDF header" (recommended)',
            "Optionally choose a distributor override if the header match is wrong or missing",
            "Distributors without a configured template show \"Template required\" and may be disabled",
            'If override is selected and template is missing, use the configure link to set one up',
        ]
    )
    m.heading("5.4.2 Excel or Bulk Upload (Multiple Files / ZIP)", 3)
    m.bullets(
        [
            "Excel files are matched by filename (code or name)",
            "PDFs are matched from the report header, with filename as fallback",
            "No manual distributor selection is required unless you want one distributor for every file",
            'Optional: check "Apply same distributor to all files" to force one distributor for every file',
            "When an Excel-only distributor is selected, the drop zone switches to Excel or ZIP only",
        ]
    )
    m.heading("5.5 Upload Steps", 2)
    m.steps(
        [
            "Click Upload in the sidebar navigation",
            "Drag and drop PDF, Excel, or ZIP files onto the upload area, or click Browse Files",
            "Remove unwanted files using the X button next to each file name",
            'Set the report date ("Business date this file represents") — cannot be in the future',
            "Optionally set a distributor override (single PDF) or apply-to-all (bulk / Excel)",
            'Click "Upload Document(s)" to start processing',
            "Watch the Processing Pipeline stepper on the right: Upload → Extract → Match Products → Review Exceptions",
            "For a single successful upload, you are automatically redirected to the document detail page",
            "For bulk uploads, review the Upload Results card and click Open document for each file",
        ]
    )
    m.heading("5.6 What Happens During Processing", 2)
    m.paragraph(
        "When you upload, the system performs these steps automatically for each file:"
    )
    m.steps(
        [
            "Detects or applies the distributor and loads its PDF template or Excel column map",
            "Enforces input mode (Excel-only distributors reject PDFs)",
            "Saves the uploaded file to secure storage",
            "For PDFs: sends the file to the extraction service (PDF worker on port 8000)",
            "For Excel: parses the workbook using the distributor's Excel column map",
            "Extracts product rows (table / line parser / alternate settings for PDF; mapped columns for Excel)",
            "Matches extracted rows against the product catalog using fuzzy matching",
            "Saves extracted rows and match results to the database",
            "Sets document status based on extraction and match results",
        ]
    )
    m.heading("5.7 Extraction Methods", 2)
    m.paragraph(
        "After processing, a toast notification may show which extraction method was used. "
        "The system tries PDF methods automatically — no user action is required."
    )
    m.table(
        ["Method", "Display Label", "When Used"],
        [
            ["table", "Table extraction", "Default — PDF has a parseable table grid"],
            ["line_fallback", "Line parser", "Fallback when table extraction yields poor results"],
            ["alternate_settings", "Alternate pdfplumber", "Retry with alternate pdfplumber parsing settings"],
            ["excel", "Excel", "Sales workbook parsed via the distributor Excel column map"],
        ],
    )
    m.heading("5.8 Document Status After Upload", 2)
    m.table(
        ["Status", "Meaning", "Your Next Action"],
        [
            ["EXTRACTED", "All rows matched successfully", "Approve & Promote to Sales"],
            ["REVIEW_REQUIRED", "Some rows could not be matched", "Resolve in Review Queue"],
            ["TEMPLATE_MISMATCH", "PDF layout differs from saved template", "Re-map PDF Template on document detail"],
            ["FAILED", "Extraction error, Excel-only block, or distributor not detected", "Check error message; retry upload or fix template / Excel map"],
            ["PROCESSING", "Upload is still being processed", "Wait for completion"],
        ],
    )
    m.heading("5.9 Upload Results Panel", 2)
    m.paragraph(
        "When uploading multiple files, an Upload Results card appears listing each file with:"
    )
    m.bullets(
        [
            "File name and status badge (EXTRACTED, REVIEW_REQUIRED, TEMPLATE_MISMATCH, FAILED)",
            "Matched distributor name (or suggested distributor if detection failed)",
            "Error message for failed files",
            "Open document link for successfully processed files",
        ]
    )
    m.heading("5.10 Warnings and Notifications", 2)
    m.bullets(
        [
            "Duplicate warning if an approved document already exists for the same distributor and report date",
            "Template mismatch toast when one or more documents need template re-mapping",
            "Extraction method toast showing table, line parser, alternate pdfplumber, or excel usage",
            "Excel map required / Template required messaging on the distributor selector",
            "Per-file warnings displayed as error toasts for partial batch failures",
        ]
    )

    # --- 6. Document Management ---
    m.heading("6. Document Management", 1)
    m.heading("6.1 Documents List", 2)
    m.paragraph(
        "The Documents page (/documents) shows all uploaded reports (PDF and Excel) with their "
        "current status, distributor, report date, period, match counts, and upload time."
    )
    m.bullets(
        [
            "Click any document row to open its detail page",
            "Use status filters to find documents needing attention",
            "Sort by date to see the most recent uploads first",
        ]
    )
    m.heading("6.2 Document Detail Page", 2)
    m.paragraph(
        "The document detail page (/documents/[id]) provides full information about a "
        "single uploaded report."
    )
    m.heading("6.2.1 Page Sections", 3)
    m.table(
        ["Section", "Description"],
        [
            ["Header", "File name, distributor, report date, status badge, and action buttons"],
            ["Processing Timeline", "Stepper showing Uploaded → Extracted → Product Match → Review → Approved"],
            ["Document Details", "Report date, distributor, file size, uploaded time, match rate, extract method, errors"],
            ["Extracted Rows Table", "Raw product text, matched product, Sales Units, S.P, Sales Value, optional Returns Qty / Closing Stock"],
        ],
    )
    m.heading("6.2.2 Available Actions", 3)
    m.table(
        ["Action", "When Available", "Result"],
        [
            ["Approve & Promote to Sales", "Status is EXTRACTED with no unresolved rows", "Creates DailySalesFact records; status becomes APPROVED"],
            ["Resolve in Review Queue", "Status is REVIEW_REQUIRED", "Opens Review Queue to resolve unmatched rows"],
            ["Re-map PDF Template", "Status is TEMPLATE_MISMATCH", "Opens Template Wizard to fix column mapping"],
            ["Retry Upload", "Status is FAILED", "Returns to Upload page to submit a new file"],
            ["Generate SSR", "Status is APPROVED", "Creates an SSR report for this document's report date"],
        ],
    )
    m.heading("6.3 Document Status Lifecycle", 2)
    m.paragraph("Documents progress through the following statuses:")
    m.paragraph(
        "UPLOADED → PROCESSING → EXTRACTED | REVIEW_REQUIRED | TEMPLATE_MISMATCH | FAILED"
    )
    m.paragraph("After review (if needed): EXTRACTED → APPROVED → SSR Generated")

    # --- 7. Review Queue ---
    m.heading("7. Review Queue", 1)
    m.paragraph(
        "The Review Queue (/review) is where you resolve product rows that could not be "
        "automatically matched to the catalog. A badge on the sidebar shows the count of "
        "pending review items."
    )
    m.heading("7.1 When Review Is Required", 2)
    m.bullets(
        [
            "Product name in the PDF does not match any catalog entry",
            "Multiple possible matches exist and confidence is too low",
            "Product is new and not yet in the catalog",
            "Distributor uses a non-standard product name or abbreviation",
        ]
    )
    m.heading("7.2 Resolving Review Items", 2)
    m.steps(
        [
            "Navigate to Review Queue from the sidebar (note the badge count)",
            "Review each unmatched row showing the extracted product name and suggested matches",
            "Choose one of the resolution options for each row (see table below)",
            "After resolving all rows, return to the document detail page",
            "Document status will update to EXTRACTED when all rows are resolved",
            "Proceed to Approve & Promote to Sales",
        ]
    )
    m.heading("7.3 Resolution Options", 2)
    m.table(
        ["Option", "When to Use", "Effect"],
        [
            ["Approve", "System suggestion is correct", "Maps row to suggested catalog product; saves mapping for future uploads"],
            ["Map", "Correct product selected from dropdown", "Maps row to selected product; saves mapping for future uploads"],
            ["New Product", "Product is genuinely new", "Creates catalog entry (Group required) and maps the row"],
            ["Reject", "Row is invalid, duplicate, or not a real product", "Excludes row from approved sales data"],
        ],
    )
    m.paragraph(
        "When creating a new product from Review, you must select a product group "
        "(from active groups on Product Groups — typically Medicronis or Transformer). "
        "Quantity on review cards is labeled Sales Units to match SSR terminology."
    )
    m.paragraph(
        "Important: Approved mappings are saved as DistributorProductMapping records. "
        "Future uploads from the same distributor will automatically match these product names."
    )
    m.heading("7.4 Example Workflow with Review", 2)
    m.paragraph(
        "Sample file: pharmalink_jul2026.pdf (in samples/pdfs/ after running the sample generator)"
    )
    m.steps(
        [
            "Upload pharmalink_jul2026.pdf with auto-detect enabled",
            "Document status becomes REVIEW_REQUIRED",
            "Open Review Queue and resolve each exception",
            "Return to document detail — status becomes EXTRACTED",
            'Click "Approve & Promote to Sales"',
            "Generate SSR and download Excel",
        ]
    )

    # --- 8. SSR Reports ---
    m.heading("8. SSR Reports", 1)
    m.paragraph(
        "SSR (Secondary Sales Report) reports consolidate approved sales data into "
        "Medicronis-format Excel workbooks. The on-screen preview and downloaded file "
        "both use the same DATA sheet column layout."
    )
    m.heading("8.1 Generating Reports from the Reports Page", 2)
    m.steps(
        [
            "Navigate to SSR Reports (/reports) from the sidebar",
            "Select the report view: Day, Week, or Month",
            'Set the "As-of date" for the reporting period (cannot be in the future)',
            "Review the coverage indicator showing how many distributors have data vs active",
            'Click "Generate SSR" to create the consolidated report',
            "You will be redirected to the report detail page",
        ]
    )
    m.heading("8.2 Generating Reports from a Document", 2)
    m.steps(
        [
            "Open an approved document's detail page",
            'Click "Generate SSR"',
            "A day-view SSR report is created for that document's report date",
        ]
    )
    m.heading("8.3 Report Detail and Download", 2)
    m.paragraph(
        "The report detail page (/reports/[id]) shows a full DATA-sheet preview titled "
        '"Updated Sales Till DD-MM-YY", with the same columns as the Excel download. '
        "Product Name and Sales Value are emphasized for quick scanning."
    )
    m.steps(
        [
            "Open the report from the Reports list or after generation",
            "Scroll horizontally through the preview grid to review all metrics",
            'Click "Download Excel" to save the .xlsx file (DATA sheet)',
            "Open the downloaded file in Microsoft Excel or compatible spreadsheet software",
        ]
    )
    m.heading("8.4 Report Views and Date Ranges", 2)
    m.table(
        ["View", "Period Covered", "Prior Comparison (\"Yesterday\" columns)"],
        [
            ["Day", "Single as-of date", "Previous calendar day"],
            ["Week", "Week containing the as-of date", "Prior week window"],
            ["Month", "Month containing the as-of date", "Prior month matching window"],
        ],
    )
    m.paragraph(
        "LMTD (Last Month To Date) columns always compare against the same calendar day "
        "in the prior month, regardless of Day/Week/Month view."
    )
    m.heading("8.5 SSR DATA Sheet Columns", 2)
    m.paragraph(
        "Each row is one distributor × product combination from the active master grid. "
        "Products without sales in the period still appear with zero units where applicable."
    )
    m.table(
        ["Column", "Description"],
        [
            ["Distributor Name, City, Region, Country", "Distributor master data"],
            ["Category", "Always \"Distributor\" for this grid"],
            ["Group", "Product group from catalog (managed on Product Groups)"],
            ["Manager", "Assigned sales manager for the distributor"],
            ["Product Name", "Catalog product name"],
            ["S.P", "Selling price used for value calculations"],
            ["Sales Units", "Units sold in the selected period"],
            ["Closing Stock", "Latest closing stock on/before as-of date (— if missing)"],
            ["Sales Value", "Sales Units × S.P"],
            ["Stock Value", "Closing Stock × S.P (0 when stock missing in Excel)"],
            ["Yesterday / Yesterday Sale Value", "Prior-period units and value (label kept for Medicronis workbook parity)"],
            ["Difference", "Current Sales Value − prior-period sale value"],
            ["LMTD Sales Unit / Value", "Same-day-prior-month units and value"],
            ["LMTD Difference (units / value)", "Current − LMTD for units and value"],
            ["LMTD %age", "(Sales Value / LMTD Sales Value − 1), or \"-\" if LMTD value is zero"],
            ["Inventory", "Sales Units × 1.5 (target stock heuristic)"],
            ["Order / Order Value", "Units (and value) needed when Inventory > Closing Stock"],
            ["Excess Stock / Excess Stock Value", "Units (and value) when Closing Stock > Inventory"],
            ["Inventory Value", "Inventory × S.P"],
        ],
    )
    m.heading("8.6 Related Export Layouts", 2)
    m.paragraph(
        "The primary user download from SSR Reports is the DATA sheet workbook described above. "
        "The platform also supports a Wholeseller (Direct Party) line layout used in related "
        "exports: Party Name, Product, Quatity, S.P, Value, with Total Direct Sale and Total Sales rows."
    )

    # --- 9. Distributors ---
    m.heading("9. Managing Distributors", 1)
    m.paragraph(
        "The Distributors page (/distributors) lets you view and manage your distributor network, "
        "including input mode (PDF + Excel vs Excel only) and template readiness."
    )
    m.heading("9.1 Distributor List", 2)
    m.bullets(
        [
            "View all registered distributors with code, name, region, country, city, manager, documents, mappings, template status, input mode, and status",
            "Search by code, name, city, manager, region, or country",
            "Filter visually by Active / Inactive status badges",
            "Template badges: Ready | Template required | Excel ready | Excel map required",
            "Input mode badge: Both (PDF + Excel) or Excel only",
            "Document and mapping counts per distributor",
            "Per-row actions: configure PDF template, configure Excel map, edit, deactivate / reactivate",
        ]
    )
    m.heading("9.2 Input Modes", 2)
    m.table(
        ["Input mode", "Allowed uploads", "Required configuration"],
        [
            ["Both (PDF + Excel)", "PDF, Excel, ZIP (mixed)", "PDF template required for PDF; Excel map required before Excel upload"],
            ["Excel only", "Excel or ZIP of Excel only (PDFs blocked)", "Excel column map required"],
        ],
    )
    m.paragraph(
        "Default is Both. Set Excel only only when you intentionally want to disallow PDF. "
        "After create, new distributors redirect to the PDF template wizard or Excel map wizard based on input mode."
    )
    m.heading("9.3 Adding a Distributor", 2)
    m.steps(
        [
            "Navigate to Distributors",
            'Click "Add Distributor"',
            "Enter distributor details: code, name, region, country, city, input mode, and assigned manager",
            "Save the new distributor record",
            "Complete the PDF template wizard (Both mode) or Excel map wizard (Excel only) when redirected",
            "Optionally configure the other map later from the distributor row actions",
        ]
    )
    m.heading("9.4 Bulk Upload", 2)
    m.steps(
        [
            'On the Distributors page, click "Download Template" if needed',
            "Fill in distributor details following the template columns",
            'Click "Bulk Upload" and submit the completed .xlsx file',
            "Review imported records and configure PDF templates / Excel maps as needed",
        ]
    )
    m.heading("9.5 Distributor Fields", 2)
    m.table(
        ["Field", "Description", "Required"],
        [
            ["Code", "Unique distributor identifier (also used in Excel filenames)", "Yes"],
            ["Name", "Full distributor business name (used for PDF header matching)", "Yes"],
            ["Region", "Geographic region", "Yes"],
            ["Country", "Country of operation", "Yes"],
            ["City", "City location", "Optional"],
            ["Input mode", "Both (PDF + Excel) or Excel only", "Yes"],
            ["Manager", "Assigned sales manager", "Optional"],
            ["Active", "Whether distributor is active for uploads", "Yes"],
        ],
    )

    # --- 10. PDF Template Configuration ---
    m.heading("10. PDF Template Configuration", 1)
    m.paragraph(
        "Each distributor's PDF reports have a unique layout. Before uploading documents "
        "for a distributor, you must configure a PDF template that maps the PDF columns "
        "to the system's data fields."
    )
    m.heading("10.1 When Template Configuration Is Required", 2)
    m.bullets(
        [
            "Adding a new distributor in Both (PDF + Excel) mode",
            "Distributor's PDF layout has changed",
            "Upload results in TEMPLATE_MISMATCH status",
            "Template required badge appears on the Upload, Dashboard Template Coverage, or Distributors page",
        ]
    )
    m.paragraph(
        "Excel-only distributors skip the PDF template wizard and use the Excel column map instead (Section 11). "
        "Opening /distributors/[id]/template for an Excel-only distributor redirects to the Excel map wizard."
    )
    m.heading("10.2 Template Mismatch Triggers", 2)
    m.paragraph(
        "The system flags TEMPLATE_MISMATCH when any of the following occur:"
    )
    m.bullets(
        [
            "PDF layout cannot be parsed with the saved template mapping",
            "Column headers cannot be resolved from the saved template mapping",
            "Extracted row count drops more than 30% compared to the last successful upload",
        ]
    )
    m.heading("10.3 Template Wizard Steps", 2)
    m.paragraph(
        "Access the template wizard from Distributors → select distributor → "
        "Configure PDF Template (/distributors/[id]/template), or from the "
        "Re-map PDF Template button on a mismatched document."
    )
    m.steps(
        [
            "Upload a sample PDF from the distributor",
            "Review the auto-suggested PDF format family (A through J)",
            "Map header columns to system fields (product name, sales units / quantity, net sale, closing stock, etc.)",
            "Preview extraction results to verify row data is correct",
            "Adjust advanced settings if needed (skip rows, extraction mode, line-fallback)",
            "For Family J (line parser), set Sales Units and Sales Value column indices until the live preview looks correct",
            "Save the template configuration",
            "Return to Upload and process PDFs for this distributor",
        ]
    )
    m.heading("10.4 PDF Format Families", 2)
    m.paragraph(
        "The system supports 10 PDF layout families (A–J) covering 46 real distributor "
        "report formats from July Closing analysis. Each family has a preset column structure."
    )
    m.table(
        ["Family", "Description", "Example Distributors"],
        [
            ["A", "SSR Stock & Return — grouped two-row header", "AIM Pharma HYD, Reliable Pharma MPK"],
            ["B", "Medicronis ERP Sales and Stock Statement", "AQ Enterprises Sargodha, N&G Enterprises SWL"],
            ["C", "Item Description / Net Sales", "Bashir Pharma Gujranwala, M&C RWP"],
            ["D", "Product Name T.P. Tax — 19-column", "Al Makkah Trading Buner, Umar Medicine Swabi"],
            ["E", "TRAD RATE / NET SALE", "Hamed Pharma Bannu, Mehran Traders Mardan"],
            ["F", "Simple Description / Sales Qty / Sale Value", "Simple single-row header layouts"],
            ["G", "Code / Product / Net Sale (label-path + column fallbacks)", "Chishti Pharma Multan, Zavion Pharma Jhang"],
            ["H", "NAME / PRICE / OPEN STOCK / SALES — 21-col grouped", "Mashal Enterprises Bajaur"],
            ["I", "Vertical NET SALES QTY.BON AMOUNT", "Hashmani Health Care Karachi"],
            ["J", "No table — line parser only (table extraction disabled)", "Ayan Pharma Taunsa, Globar Enterprises SKP, others"],
        ],
    )
    m.paragraph(
        "Family J PDFs have no reliable table grid. The Template Wizard enables line-parser "
        "settings so you can map which numeric token positions are Sales Units and Sales Value. "
        "Scanned PDFs without extractable text still require OCR (Tesseract) on the worker. "
        "See docs/pdf-families.md for the full distributor list."
    )
    m.heading("10.5 Field Mapping Reference", 2)
    m.table(
        ["System Field", "Typical PDF Column Names", "Required"],
        [
            ["Product Name", "ITEM, DESCRIPTION, Product Name, Item Description", "Yes"],
            ["Quantity / Sales Units", "QTY, SALE, Sales Qty, Quantity Sold", "Yes"],
            ["Net Sale", "NET SALE, Net Sales, Sale Value, AMOUNT", "Yes"],
            ["Closing Stock", "CLOSING, Closing Balance, Stock", "Optional"],
            ["Rate / TP", "RATE, TP, TRAD RATE, T.P.", "Optional"],
        ],
    )

    # --- 11. Excel Column Map ---
    m.heading("11. Excel Column Map Configuration", 1)
    m.paragraph(
        "Each distributor that accepts Excel sales workbooks needs an Excel column map. "
        "The map tells the system which workbook columns hold product name, sales quantity, "
        "unit price, closing stock, and related fields."
    )
    m.heading("11.1 When Excel Map Configuration Is Required", 2)
    m.bullets(
        [
            "Distributor input mode is Excel only (required before any upload)",
            "Distributor input mode is Both and you plan to upload .xlsx sales files",
            "Excel map required badge appears on Distributors, Upload, or Template Coverage",
            "Adding a new Excel-only distributor (wizard opens automatically after create)",
        ]
    )
    m.heading("11.2 Excel Map Wizard Steps", 2)
    m.paragraph(
        "Access the wizard from Distributors → row actions → Excel map "
        "(/distributors/[id]/excel-template), or after creating an Excel-only distributor."
    )
    m.steps(
        [
            "Open Excel map for the distributor (title: \"Excel map — [name]\")",
            "Set Header row (0-based) and optional Sheet name (first sheet if blank)",
            "Upload a sample .xlsx workbook and refresh the live preview if needed",
            "Map columns: Product name* and Sales qty* are required",
            "Optionally map Unit price (S.P), Closing stock (recommended), Sales amount, Returns qty",
            "Review the Live preview table (Product, Qty, Unit price, Closing, Amount)",
            'Click "Save Excel map"',
            "Confirm the badge shows Excel map ready / Excel ready on Distributors",
        ]
    )
    m.heading("11.3 Excel Field Mapping Reference", 2)
    m.table(
        ["Field", "Required", "Notes"],
        [
            ["Product name", "Yes", "Raw product text used for catalog matching"],
            ["Sales qty", "Yes", "Sales Units promoted to daily sales facts"],
            ["Unit price (S.P)", "Recommended", "Selling price for value calculations"],
            ["Closing stock", "Recommended", "Used on SSR Closing Stock / Stock Value"],
            ["Sales amount", "Optional", "Line amount when provided by the workbook"],
            ["Returns qty", "Optional", "Returns quantity when present"],
        ],
    )
    m.paragraph(
        "Distributors in Both mode can keep a PDF template and an Excel map. Use the "
        "\"PDF template\" link on the Excel map page (or vice versa) to switch between wizards."
    )

    # --- 12. Products ---
    m.heading("12. Managing Products", 1)
    m.paragraph(
        "The Products page (/products) manages the product catalog used for matching "
        "extracted PDF/Excel rows and for SSR DATA sheet line items (including Group)."
    )
    m.heading("12.1 Product Catalog", 2)
    m.bullets(
        [
            "View all products with SKU, name, group, manufacturer, category, and bonus",
            "Search by SKU, name, manufacturer, group, category, or aliases",
            "Each product must belong to an active product group (see Section 13)",
            "Aliases improve fuzzy matching from distributor PDF/Excel text",
            "Actions: Add, Edit, View, Deactivate; Download Template and Bulk Upload",
        ]
    )
    m.heading("12.2 Adding a Product", 2)
    m.steps(
        [
            "Navigate to Products",
            'Click "Add Product"',
            "Enter SKU and product name (SKU must be unique)",
            "Select Group from active product groups — required",
            "Optionally set category, composition, manufacturer (existing or new), shipper size",
            "Optionally enter pricing fields: MRP, TP, Old SP, New SP, Net Price, Tax, Net Price with 1%",
            "Optionally enter Bonus in N+N format (for example 4+1)",
            "Add comma-separated aliases if the product appears under alternate names in reports",
            "Save the product",
        ]
    )
    m.heading("12.3 Bulk Product Upload", 2)
    m.steps(
        [
            'Click "Download Template" on the Products page to get the Medicronis product Excel template',
            "Fill rows using the template columns (see table below)",
            "Set Group to an existing product group name on every row (e.g. Medicronis or Transformer)",
            "Upload the completed .xlsx file via Bulk Upload",
            "Review any row-level import errors, then verify products in the catalog",
        ]
    )
    m.heading("12.4 Product Import Template Columns", 2)
    m.table(
        ["Column", "Required", "Notes"],
        [
            ["SKU", "Yes", "Unique product code"],
            ["Product Name", "Yes", "Catalog display name"],
            ["Category", "No", "Optional classification"],
            ["Group", "Yes", "Must match an existing product group name"],
            ["Composition", "No", "Formula / strength text"],
            ["Manufacturer", "No", "Created automatically if new"],
            ["Shipper Size", "No", "Integer pack size"],
            ["MRP, TP, Old SP, New SP", "No", "Pricing decimals"],
            ["Net Price, Tax, Net Price with 1%", "No", "Pricing decimals"],
            ["Bonus", "No", "Format N+N (e.g. 10+2)"],
            ["Aliases", "No", "Alternate names for matching"],
        ],
    )
    m.heading("12.5 Product Aliases", 2)
    m.paragraph(
        "Aliases are alternate names for a product used during fuzzy matching. "
        "For example, a product might appear in distributor reports as an abbreviated name "
        "or regional spelling variant. Adding aliases improves automatic match rates."
    )

    # --- 13. Product Groups ---
    m.heading("13. Managing Product Groups", 1)
    m.paragraph(
        "The Product Groups page (/product-groups) manages the groups used on products, "
        "bulk imports, Review Queue new-product creation, and the SSR DATA sheet Group column."
    )
    m.heading("13.1 Product Groups List", 2)
    m.bullets(
        [
            "View all groups with name, active status, and product count",
            "Search by group name",
            "Seeded defaults typically include Medicronis and Transformer",
            "You can add additional groups as your portfolio grows",
        ]
    )
    m.heading("13.2 Adding or Editing a Group", 2)
    m.steps(
        [
            "Navigate to Product Groups from the sidebar",
            'Click "Add Product Group" (or Edit on an existing row)',
            "Enter a unique group name",
            "Save — the group becomes available in product forms and imports",
        ]
    )
    m.heading("13.3 Deactivate / Reactivate", 2)
    m.bullets(
        [
            "Deactivate a group that should no longer be assigned to new products",
            "Reactivate when the group is needed again",
            "Products already assigned to a deactivated group keep their assignment until edited",
        ]
    )
    m.paragraph(
        "Group is required when creating or editing a product, when bulk-importing products, "
        "and when creating a product from the Review Queue."
    )

    # --- 14. Complete Workflows ---
    m.heading("14. Complete Workflows", 1)
    m.heading("14.1 Quick Demo Path — PDF (No Review Required)", 2)
    m.paragraph(
        "Use sample file medsupply_jul2026.pdf for the fastest end-to-end demonstration."
    )
    m.steps(
        [
            "Log in with demo password: demo → Access Demo",
            "Go to Upload",
            "Set report date and upload medsupply_jul2026.pdf (auto-detect distributor from header)",
            "You are redirected to document detail — verify extracted rows with high match rate",
            'Click "Approve & Promote to Sales"',
            "Go to SSR Reports or click Generate SSR on the document",
            "On Reports: set As-of date and View (Day/Week/Month), then Generate SSR",
            'Click "Download Excel" and open the .xlsx file',
        ]
    )
    m.heading("14.2 Quick Demo Path — Excel Upload", 2)
    m.steps(
        [
            "Confirm the target distributor has Excel map ready (configure if needed — Section 11)",
            "Name the workbook with distributor code or name (e.g. AYAN-TAUNSA.xlsx)",
            "Go to Upload, set report date, drop the .xlsx file",
            "Leave auto-match by filename (or apply-to-all override)",
            "Review document detail → Approve & Promote to Sales",
            "Generate SSR (Day view from document, or Day/Week/Month from Reports)",
        ]
    )
    m.heading("14.3 Full Workflow with Review", 2)
    m.steps(
        [
            "Log in to the application",
            "Upload pharmalink_jul2026.pdf with auto-detect enabled",
            "Document status becomes REVIEW_REQUIRED",
            "Open Review Queue from sidebar (check badge count)",
            "Resolve each unmatched row (Approve, Map, New Product, or Reject)",
            "Return to document detail — status becomes EXTRACTED",
            'Click "Approve & Promote to Sales"',
            "Generate SSR report from Reports (As-of date + View) or from the document",
            "Download and verify Excel output",
        ]
    )
    m.heading("14.4 Bulk Upload Workflow", 2)
    m.steps(
        [
            "Collect PDF and/or Excel sales reports from multiple distributors",
            "Go to Upload and drag all files (or a ZIP) into the upload area",
            "Set the shared report date for all files",
            "Leave distributor on auto-detect (PDF header / Excel filename)",
            'Click "Upload Documents"',
            "Review the Upload Results card for each file's status",
            "Resolve review items, approve each document, then generate consolidated SSR",
        ]
    )
    m.heading("14.5 Setting Up a New Distributor", 2)
    m.steps(
        [
            "Add the distributor on the Distributors page (choose Input mode: Both or Excel only)",
            "Complete the PDF template wizard (Both) or Excel map wizard (Excel only)",
            "Optionally configure the other map if Both mode will use both file types",
            "Ensure catalog products have the correct product group",
            "Upload the first real sales report (PDF or Excel as allowed by input mode)",
            "Complete review if needed, then approve and generate SSR",
        ]
    )
    m.heading("14.6 Monthly Reporting Cycle", 2)
    m.steps(
        [
            "Collect PDF/Excel sales reports from all active distributors for the month",
            "Upload files individually or in bulk with the correct report date",
            "Resolve all review queue items across all documents",
            "Approve all extracted documents",
            "Go to SSR Reports, select Month view, set the month-end as-of date",
            "Generate consolidated SSR and download Excel",
            "Distribute the SSR report to stakeholders",
        ]
    )

    # --- 15. Troubleshooting ---
    m.heading("15. Troubleshooting", 1)
    m.heading("15.1 Common Issues and Solutions", 2)
    m.table(
        ["Issue", "Possible Cause", "Solution"],
        [
            ["Upload fails or 'Extraction failed'", "PDF worker service not running", "Ask administrator to start PDF worker; verify health at port 8000"],
            ["Could not detect distributor from PDF header", "Distributor name not in PDF or not registered", "Register distributor or use manual override on Upload page"],
            ["Could not match distributor from Excel filename", "Filename missing code/name", "Rename file (e.g. CODE.xlsx) or use Apply same distributor"],
            ["Excel-only / PDF blocked", "Distributor input mode is Excel only", "Upload .xlsx instead; configure Excel map if missing"],
            ["Excel map required", "No Excel column map saved", "Open Excel map wizard (Section 11) and save mapping"],
            ["Database errors on any page", "Database not initialized", "Administrator: run migrations and seed data"],
            ["Empty dashboard sales charts", "No approved documents yet", "Approve at least one document to populate charts"],
            ["Template required on upload", "Distributor has no configured PDF template", "Configure template via Template Wizard (Section 10)"],
            ["TEMPLATE_MISMATCH status", "PDF layout changed or wrong mapping", "Click Re-map PDF Template; check row count drop >30%"],
            ["Review queue won't clear", "Unresolved UNMATCHED/PENDING rows remain", "Resolve all rows in Review Queue before approval"],
            ["Approval blocked", "Unresolved review rows exist", "Complete all review items first"],
            ["Duplicate upload warning", "Approved doc exists for same distributor + date", "Expected; only re-upload to replace existing data"],
            ["Scanned PDF fails", "No extractable text in PDF", "Administrator: install Tesseract OCR on worker"],
            ["Downloaded Excel is empty", "No approved data for selected as-of period", "Approve documents for that date first; check Day/Week/Month view"],
            ["Group column blank in SSR", "Product missing product group", "Edit product and assign an active group (Section 13)"],
            ["Product save / import fails on Group", "Missing or unknown group name", "Create the group on Product Groups, then use that exact name"],
            ["Family J extraction wrong quantities", "Line-parser column indices incorrect", "Re-open Template Wizard; adjust Sales Units / Sales Value indices"],
            ["Session expired", "7-day cookie expired", "Log in again with Access Demo"],
        ],
    )
    m.heading("15.2 Document Status Reference", 2)
    m.table(
        ["Status", "Description"],
        [
            ["UPLOADED", "File saved, processing not yet started"],
            ["PROCESSING", "Extraction and matching in progress"],
            ["EXTRACTED", "All rows extracted and matched — ready for approval"],
            ["REVIEW_REQUIRED", "Some rows need manual resolution in Review Queue"],
            ["TEMPLATE_MISMATCH", "PDF layout does not match configured template"],
            ["FAILED", "Processing error — check file type, Excel-only rules, PDF quality, and retry"],
            ["APPROVED", "Data promoted to sales facts — ready for SSR generation"],
        ],
    )
    m.heading("15.3 Getting Help", 2)
    m.bullets(
        [
            "Contact your system administrator for login issues, service outages, or template configuration",
            "Refer to docs/pdf-families.md for detailed PDF format family documentation",
            "Refer to docs/excel-only-distributors.md for historical Excel-only notes (default is Both)",
            "Run the smoke test checklist (SMOKE_TEST.md) to validate system health",
        ]
    )

    # --- 16. Appendix ---
    m.heading("16. Appendix", 1)
    m.heading("16.1 Navigation Reference", 2)
    m.table(
        ["Page", "Route", "Purpose"],
        [
            ["Dashboard", "/dashboard", "KPIs, sales charts, template coverage, activity"],
            ["Upload", "/upload", "Submit PDF, Excel, or ZIP sales reports"],
            ["Documents", "/documents", "View all uploaded documents"],
            ["Document Detail", "/documents/[id]", "View rows, approve, generate SSR"],
            ["Review Queue", "/review", "Resolve unmatched product rows"],
            ["SSR Reports", "/reports", "Generate by as-of date + Day/Week/Month"],
            ["Report Detail", "/reports/[id]", "Preview DATA sheet and download Excel"],
            ["Distributors", "/distributors", "Manage network, input mode, templates"],
            ["PDF Template Wizard", "/distributors/[id]/template", "Configure PDF column mapping"],
            ["Excel Map Wizard", "/distributors/[id]/excel-template", "Configure Excel column map"],
            ["Products", "/products", "Manage product catalog"],
            ["Product Groups", "/product-groups", "Manage groups for products and SSR"],
            ["Login", "/login", "Demo password authentication"],
        ],
    )
    m.heading("16.2 Glossary", 2)
    m.table(
        ["Term", "Definition"],
        [
            ["SSR", "Secondary Sales Report — consolidated Medicronis Excel export of distributor sales data"],
            ["DATA sheet", "Primary SSR workbook sheet with distributor × product metrics, LMTD, and inventory columns"],
            ["As-of date", "Business date used when generating Day/Week/Month SSR aggregations"],
            ["Auto-detect", "Automatic distributor identification from PDF header text or Excel filename"],
            ["Input mode", "Distributor setting: Both (PDF + Excel) or Excel only"],
            ["Excel map", "Per-distributor column mapping for .xlsx sales workbooks"],
            ["PDF Format Family", "Preset layout template (A–J) for a class of distributor PDF reports"],
            ["Family J / Line parser", "Text-line extraction mode when the PDF has no usable table grid"],
            ["Template Config", "Per-distributor PDF column mapping stored in the database"],
            ["Product Group", "Catalog attribute managed at /product-groups — appears as Group on SSR"],
            ["Daily Sales Fact", "Approved sales record promoted from an extracted document row"],
            ["Match Rate", "Percentage of extracted rows successfully matched to catalog products"],
            ["Review Queue", "List of unmatched rows requiring manual resolution"],
            ["Distributor Product Mapping", "Saved alias linking a distributor's product name to a catalog SKU"],
            ["LMTD", "Last Month To Date — same calendar day in the prior month for comparison metrics"],
            ["Inventory (SSR)", "Computed target stock = Sales Units × 1.5"],
            ["Processing Pipeline", "Upload → Extract → Match Products → Review Exceptions stepper on Upload page"],
            ["S.P", "Selling price used for Sales Value and related SSR calculations"],
        ],
    )
    m.heading("16.3 Known Limitations (Demo Version)", 2)
    m.bullets(
        [
            "Single shared password — no multi-user role-based access control",
            "Synchronous processing — no background job queue",
            "10 PDF format families — custom parsers require further development",
            "OCR is optional — scanned PDFs may fail without Tesseract installed",
            "No email notifications, audit trail, or ERP integration",
            "Cloud deployments may use ephemeral file storage (files lost on cold start)",
            "Trial banner and \"Trial — Medicronis\" badge indicate the shared demo environment",
        ]
    )
    m.heading("16.4 Sample Files", 2)
    m.paragraph("Sample PDF files for testing can be generated with:")
    m.code_line("python scripts/generate_sample_pdfs.py")
    m.bullets(
        [
            "medsupply_jul2026.pdf — quick demo path, no review required",
            "pharmalink_jul2026.pdf — demo with review queue exceptions",
            "Excel demos: use a real distributor workbook named with code/name after saving an Excel map",
        ]
    )

    m.page_break()
    footer_p = m.doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run(
        f"Medicronis Sales Intelligence User Manual — Version {MANUAL_VERSION} — "
        f"{date.today().strftime('%B %Y')}"
    )
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(100, 100, 100)
    run.italic = True


def build_manual() -> Path:
    doc = Document()
    title_break = add_title_page(doc)

    m = ManualBuilder(doc=doc)
    write_content(m)

    toc_elements = build_toc_elements(doc, m.toc)
    anchor = title_break
    for el in toc_elements:
        anchor.addnext(el)
        anchor = el

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build_manual()
    print(f"User manual created: {path}")
