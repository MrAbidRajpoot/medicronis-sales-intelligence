"""Generate CEO meeting documents: PPTX, DOCX, and PDF."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUTPUT_DIR = Path(__file__).resolve().parent.parent / "docs" / "ceo-meeting"
MEETING_DATE = "[Meeting Date]"
TRIAL_DATES = "[Start Date] to [End Date]"

# Brand colors
PRIMARY = RGBColor(0x1A, 0x56, 0x8E)  # Deep blue
ACCENT = RGBColor(0x0E, 0x7A, 0x6B)  # Teal
DARK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)

PPT_PRIMARY = PptRGBColor(0x1A, 0x56, 0x8E)
PPT_ACCENT = PptRGBColor(0x0E, 0x7A, 0x6B)
PPT_DARK = PptRGBColor(0x1E, 0x29, 0x3B)
PPT_MUTED = PptRGBColor(0x64, 0x74, 0x8B)
PPT_WHITE = PptRGBColor(0xFF, 0xFF, 0xFF)


def ensure_output_dir() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR


# ---------------------------------------------------------------------------
# PowerPoint helpers
# ---------------------------------------------------------------------------


def set_slide_bg(slide, color: PptRGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, PPT_PRIMARY)

    box = slide.shapes.add_textbox(PptInches(0.6), PptInches(2.0), PptInches(8.8), PptInches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = PptPt(36)
    p.font.bold = True
    p.font.color.rgb = PPT_WHITE
    p.alignment = PP_ALIGN.LEFT

    sub = slide.shapes.add_textbox(PptInches(0.6), PptInches(3.6), PptInches(8.8), PptInches(1.2))
    stf = sub.text_frame
    sp = stf.paragraphs[0]
    sp.text = subtitle
    sp.font.size = PptPt(18)
    sp.font.color.rgb = PPT_WHITE
    sp.alignment = PP_ALIGN.LEFT

    footer = slide.shapes.add_textbox(PptInches(0.6), PptInches(6.8), PptInches(8.8), PptInches(0.5))
    ftf = footer.text_frame
    fp = ftf.paragraphs[0]
    fp.text = f"Prepared for: Medicronis  |  Prepared by: MITRIXS — Abid Ali  |  {MEETING_DATE}"
    fp.font.size = PptPt(11)
    fp.font.color.rgb = PPT_WHITE


def add_content_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    subtitle: str | None = None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, PPT_WHITE)

    # Title bar
    bar = slide.shapes.add_shape(1, PptInches(0), PptInches(0), PptInches(10), PptInches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PPT_PRIMARY
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(PptInches(0.5), PptInches(0.15), PptInches(9), PptInches(0.7))
    ttf = title_box.text_frame
    tp = ttf.paragraphs[0]
    tp.text = title
    tp.font.size = PptPt(24)
    tp.font.bold = True
    tp.font.color.rgb = PPT_WHITE

    y = 1.3
    if subtitle:
        sub_box = slide.shapes.add_textbox(PptInches(0.6), PptInches(y), PptInches(8.8), PptInches(0.5))
        stf = sub_box.text_frame
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.font.size = PptPt(14)
        sp.font.bold = True
        sp.font.color.rgb = PPT_ACCENT
        y += 0.55

    body = slide.shapes.add_textbox(PptInches(0.6), PptInches(y), PptInches(8.8), PptInches(5.5))
    btf = body.text_frame
    btf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = PptPt(16)
        p.font.color.rgb = PPT_DARK
        p.space_after = PptPt(8)


def add_table_slide(prs: Presentation, title: str, headers: list[str], rows: list[list[str]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, PPT_WHITE)

    bar = slide.shapes.add_shape(1, PptInches(0), PptInches(0), PptInches(10), PptInches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PPT_PRIMARY
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(PptInches(0.5), PptInches(0.15), PptInches(9), PptInches(0.7))
    ttf = title_box.text_frame
    tp = ttf.paragraphs[0]
    tp.text = title
    tp.font.size = PptPt(24)
    tp.font.bold = True
    tp.font.color.rgb = PPT_WHITE

    cols = len(headers)
    table_shape = slide.shapes.add_table(len(rows) + 1, cols, PptInches(0.5), PptInches(1.4), PptInches(9), PptInches(0.4 * (len(rows) + 2)))
    table = table_shape.table

    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = header
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = PptPt(12)
            paragraph.font.color.rgb = PPT_WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = PPT_PRIMARY

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = PptPt(11)
                paragraph.font.color.rgb = PPT_DARK


def create_presentation() -> Path:
    prs = Presentation()
    prs.slide_width = PptInches(10)
    prs.slide_height = PptInches(7.5)

    add_title_slide(
        prs,
        "Automated Distributor Sales Reporting",
        "From PDF chaos to accurate SSR in minutes",
    )

    add_content_slide(
        prs,
        "The Problem Today",
        [
            "48+ distributors send sales data as PDF files in inconsistent formats",
            "One employee spends all day, every day reviewing PDFs and entering data manually",
            "SSR (Sales Stock Report) and management reports can only be produced after manual processing",
            "",
            "Impact:",
            "• Slow reporting → delayed management decisions",
            "• High risk of human error (quantities, product mapping, duplicates)",
            "• No centralized view of distributor sales performance",
        ],
    )

    add_table_slide(
        prs,
        "What We Demonstrated (Trial Proof)",
        ["Metric", "Before", "With System (Trial)"],
        [
            ["PDF processing", "Manual, all day", "[X] minutes per batch"],
            ["SSR generation", "After manual entry", "Automated after upload"],
            ["Product matching", "Manual lookup", "[X]% auto-matched; exceptions reviewed"],
            ["Formats tested", "—", "[X] of [Y] PDFs processed successfully"],
        ],
    )

    add_content_slide(
        prs,
        "Proposed Solution — Phase 1",
        [
            "A Distributor Sales Intelligence Platform — not a rip-and-replace ERP",
            "",
            "Distributor PDFs → Extract & normalize → Product matching → SSR & reports",
            "                              ↓",
            "                    Human review for exceptions only",
            "",
            "Phase 1 delivers:",
            "• Upload distributor PDFs (single or bulk)",
            "• Automatic data extraction (digital + scanned via OCR)",
            "• Map distributor products to your product master",
            "• Generate SSR in your existing format",
            "• Export Excel/PDF; basic sales visibility",
            "",
            "Works alongside your current sales software — fills the gap it cannot cover.",
        ],
    )

    add_table_slide(
        prs,
        "Why This Approach Works",
        ["Challenge", "Our Approach"],
        [
            ["48 different PDF formats", "Template rules per distributor; roll out in waves"],
            ["Same product, different names", "Fuzzy + rule-based matching + review queue"],
            ["Scanned PDFs", "OCR pipeline for non-digital documents"],
            ["Trust & accuracy", "Human validates exceptions; system learns mappings"],
            ["Existing systems", "Integration-first — no forced migration in Phase 1"],
        ],
    )

    add_content_slide(
        prs,
        "Phase 1 Scope (What You Approve Today)",
        [
            "In scope — MVP (6–8 weeks after kickoff):",
            "• User access & roles (admin, data entry, viewer)",
            "• Distributor master & document templates",
            "• Product master integration",
            "• PDF upload, processing, history",
            "• Product matching + validation screen",
            "• SSR generation (your confirmed format)",
            "• Excel/PDF export, basic dashboard, audit log",
            "",
            "Out of scope for Phase 1 (future phases):",
            "• Inventory, procurement, finance, HR, full ERP replacement",
        ],
    )

    add_content_slide(
        prs,
        "Rollout Plan",
        [
            "✓ Done — Discovery + POC + Trial demo",
            "Week 1–2 — Requirements lock, product master, top 10 distributor templates",
            "Week 3–6 — Core platform, extraction engine, SSR, review workflow",
            "Week 7–8 — UAT with your team, fix edge cases, go-live",
            "Post go-live — Add remaining distributor templates (2–4 per week)",
            "",
            "Go-live: Top 10–15 distributors (~[X]% of volume)",
            "Month 2–3: Remaining templates based on priority",
        ],
    )

    add_table_slide(
        prs,
        "Business Case / ROI",
        ["Item", "Estimate"],
        [
            ["Staff time on PDF/SSR today", "~8 hrs/day × 22 days = ~176 hrs/month"],
            ["Phase 1 target after automation", "<20 hrs/month (review exceptions only)"],
            ["Time recovered", "~150+ hrs/month"],
            ["Error reduction", "Fewer rework cycles, faster SSR turnaround"],
            ["Management benefit", "Same-day consolidated sales view vs multi-day delay"],
        ],
    )

    add_table_slide(
        prs,
        "Investment & Commercial Terms",
        ["Milestone", "Deliverable", "Payment"],
        [
            ["Kickoff", "Signed scope, access, samples", "30% — [Amount]"],
            ["Core engine live", "Upload, extract, match, review", "40% — [Amount]"],
            ["Go-live + training", "SSR production, [X] distributor templates", "30% — [Amount]"],
        ],
    )

    add_content_slide(
        prs,
        "Success Metrics (Phase 1 KPIs)",
        [
            "≥90% field extraction accuracy on digital PDFs (top templates)",
            "≥85% product auto-match rate; remainder in review queue",
            "SSR generation: same day vs current multi-day cycle",
            "Manual data entry: ≥80% reduction in hours",
            "10–15 distributor templates live at go-live",
            "System availability: ≥99% during business hours",
            "",
            "Review at 30 / 60 / 90 days post go-live.",
        ],
    )

    add_content_slide(
        prs,
        "Long-Term Vision (Optional Future Phases)",
        [
            "Phase 1 is the foundation — not the ceiling",
            "",
            "Once sales data is centralized, the same platform can grow into:",
            "• Advanced sales analytics (territory, targets, growth)",
            "• Inventory & batch/expiry (pharma-specific)",
            "• Procurement & finance integration",
            "• Single source of truth across operations",
            "",
            "Today's decision: Approve Phase 1 only.",
            "Future phases: Optional, based on proven results.",
        ],
    )

    add_content_slide(
        prs,
        "Recommended Decision & Next Steps",
        [
            "We request approval to proceed with Phase 1:",
            "",
            "1. Approve Phase 1 scope & commercial terms",
            "2. Nominate project sponsor + daily user as UAT contact",
            "3. Provide product master export + remaining PDF samples",
            "4. Kickoff within [X] days of signing",
            "",
            "Contact: Abid Ali · MITRIXS",
            "03204014920 · mitrixsofficial@gmail.com",
        ],
    )

    path = ensure_output_dir() / "CEO-Presentation.pptx"
    prs.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Word helpers
# ---------------------------------------------------------------------------


def style_doc_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = PRIMARY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def style_doc_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = PRIMARY
    p.space_before = Pt(12)


def add_body(doc: Document, text: str, bullet: bool = False) -> None:
    style = "List Bullet" if bullet else None
    p = doc.add_paragraph(text, style=style)
    for run in p.runs:
        run.font.size = Pt(11)
        run.font.color.rgb = DARK


def add_table_doc(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for r, row in enumerate(rows):
        row_cells = table.rows[r + 1].cells
        for c, value in enumerate(row):
            row_cells[c].text = value
    doc.add_paragraph()


def create_executive_one_pager() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    style_doc_title(doc, "MEDICRONIS — Automated Distributor Sales Reporting")
    style_doc_subtitle(doc, f"Executive Summary · MITRIXS · {date.today().strftime('%d %B %Y')}")

    add_section_heading(doc, "Problem")
    add_body(
        doc,
        "Medicronis receives sales reports from 48+ distributors as PDFs in inconsistent formats. "
        "Staff manually process these documents full-time before SSR and management reports can be "
        "produced — causing delays, errors, and no centralized sales view.",
    )

    add_section_heading(doc, "Solution (Phase 1)")
    add_body(
        doc,
        "MITRIXS proposes a Distributor Sales Intelligence Platform that ingests PDFs, extracts and "
        "standardizes data, maps products to your master catalog, and generates SSR automatically — "
        "with human review only for exceptions.",
    )

    add_section_heading(doc, "Proof")
    add_body(
        doc,
        f"A working demo was built on Medicronis sample PDFs and SSR format. Your team tested the "
        f"system for 2–3 days ({TRIAL_DATES}) with results:",
    )
    add_body(doc, "[Insert 2–3 bullet outcomes from trial]", bullet=True)
    add_body(doc, "[Insert user feedback quote]", bullet=True)

    add_section_heading(doc, "Phase 1 Deliverables (6–8 weeks)")
    add_body(doc, "PDF upload & processing · Product matching · Validation workflow · SSR generation · Excel/PDF export · Basic dashboard · User roles & audit trail · Go-live with top 10–15 distributor templates")

    add_section_heading(doc, "Investment")
    add_body(doc, "[Total amount] in three milestones: 30% kickoff / 40% core delivery / 30% go-live")

    add_section_heading(doc, "Expected Return")
    add_body(doc, "Recover ~150+ staff hours/month · Same-day SSR · Reduced mapping errors · Foundation for future sales analytics")

    add_section_heading(doc, "Decision Requested")
    add_body(doc, "Approve Phase 1 to begin formal development and production deployment.")

    doc.add_paragraph()
    p = doc.add_paragraph("MITRIXS — Abid Ali · 03204014920 · mitrixsofficial@gmail.com")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    path = ensure_output_dir() / "Executive-One-Pager.docx"
    doc.save(str(path))
    return path


def create_phase1_scope() -> Path:
    doc = Document()
    style_doc_title(doc, "Phase 1 — Scope of Work")
    style_doc_subtitle(doc, "Medicronis Distributor Sales Intelligence Platform · Version 1.0")

    add_section_heading(doc, "1. Background")
    add_body(doc, "Medicronis requires automation of distributor PDF sales report processing and SSR generation. Phase 1 addresses this specific operational gap without replacing existing systems.")

    add_section_heading(doc, "2. Objectives")
    objectives = [
        "Eliminate manual re-entry of distributor PDF data where possible",
        "Generate SSR in Medicronis-approved format",
        "Provide exception-based human validation",
        "Establish centralized sales data for reporting",
    ]
    for obj in objectives:
        add_body(doc, obj, bullet=True)

    add_section_heading(doc, "3. Functional Requirements")
    for heading, items in [
        ("3.1 User Management", ["Login, roles (Admin, Operator, Viewer)", "Activity logging"]),
        ("3.2 Master Data", ["Distributor profiles, document template configuration", "Product master (import from client-provided source)"]),
        ("3.3 Document Processing", ["PDF upload (single/bulk), duplicate detection", "Text/table extraction; OCR for scanned documents", "Distributor auto-identification where configured"]),
        ("3.4 Product Matching", ["Exact, code-based, and fuzzy matching", "Manual review queue; saved mappings for reuse"]),
        ("3.5 Reporting", ["SSR per agreed template", "Excel and PDF export"]),
        ("3.6 Dashboard", ["Total sales, distributor count, processing status, exceptions pending"]),
    ]:
        add_body(doc, heading)
        for item in items:
            add_body(doc, item, bullet=True)

    add_section_heading(doc, "4. Client Responsibilities")
    for item in [
        "Provide SSR template and sample outputs",
        "Provide product master data",
        "Provide representative PDFs per distributor",
        "Assign UAT contact (daily PDF user)",
        "Timely feedback during UAT (≤3 business days per review cycle)",
    ]:
        add_body(doc, item, bullet=True)

    add_section_heading(doc, "5. Out of Scope (Phase 1)")
    for item in [
        "Inventory, warehouse, procurement, finance, HR modules",
        "Replacement of existing sales/accounting software",
        "Mobile applications",
        "All 48 distributor templates at day-one go-live (rolled out post go-live)",
    ]:
        add_body(doc, item, bullet=True)

    add_section_heading(doc, "6. Timeline")
    add_table_doc(
        doc,
        ["Phase", "Duration"],
        [
            ["Kickoff & requirements lock", "Week 1"],
            ["Development", "Weeks 2–6"],
            ["UAT & fixes", "Weeks 7–8"],
            ["Go-live", "End of Week 8"],
        ],
    )

    add_section_heading(doc, "7. Acceptance Criteria")
    for item in [
        "SSR output matches agreed format for approved test PDFs",
        "Successful processing of agreed go-live distributor template set",
        "User training completed (1 session, up to [X] users)",
        "Documentation: user guide + admin guide",
    ]:
        add_body(doc, item, bullet=True)

    add_section_heading(doc, "8. Support")
    add_body(doc, "30/60/90 days warranty bug-fix period post go-live")
    add_body(doc, "Template additions for additional distributors: [X] included / [rate] per additional template")

    add_section_heading(doc, "9. Assumptions")
    for item in [
        "Client provides digital product master within 5 business days of kickoff",
        "SSR field mapping signed off in Week 1",
        "Hosting: [Client-hosted / MITRIXS-managed — specify]",
    ]:
        add_body(doc, item, bullet=True)

    path = ensure_output_dir() / "Phase-1-Scope-of-Work.docx"
    doc.save(str(path))
    return path


def create_trial_results_template() -> Path:
    doc = Document()
    style_doc_title(doc, "Medicronis Trial Results")
    style_doc_subtitle(doc, TRIAL_DATES)

    add_section_heading(doc, "Participants")
    add_body(doc, "[Names / roles]")

    add_section_heading(doc, "PDFs Tested")
    add_table_doc(
        doc,
        ["#", "Distributor", "Format Type", "Result", "Notes"],
        [["1", "", "Digital / Scanned", "✓ / ⚠ / ✗", ""]],
    )

    add_section_heading(doc, "SSR Comparison")
    add_table_doc(
        doc,
        ["Check", "Manual SSR", "System SSR", "Match?"],
        [
            ["Row count", "", "", ""],
            ["Total qty", "", "", ""],
            ["Total value", "", "", ""],
            ["Sample products (3)", "", "", ""],
        ],
    )

    add_section_heading(doc, "Time Comparison")
    add_table_doc(
        doc,
        ["Task", "Manual (est.)", "System", "Savings"],
        [
            ["Process 1 PDF", "[X] min", "[Y] min", ""],
            ["Generate SSR (batch)", "[X] hrs", "[Y] min", ""],
        ],
    )

    add_section_heading(doc, "Product Matching")
    add_body(doc, "Auto-matched: [X]%", bullet=True)
    add_body(doc, "Reviewed & corrected: [X]%", bullet=True)
    add_body(doc, "Unknown (new mapping saved): [X] products", bullet=True)

    add_section_heading(doc, "User Feedback")
    for label in ["What worked:", "What needs improvement:", "Would they use this daily? Y / N", "Quote for CEO slide:"]:
        add_body(doc, label)
        doc.add_paragraph()

    add_section_heading(doc, "Issues Found")
    add_body(doc, "1.", bullet=True)
    add_body(doc, "2.", bullet=True)
    add_body(doc, "3.", bullet=True)

    add_section_heading(doc, "Fix Plan for Phase 1")
    add_body(doc, "[How each issue will be addressed in MVP]")

    path = ensure_output_dir() / "Trial-Results-Template.docx"
    doc.save(str(path))
    return path


def create_meeting_agenda() -> Path:
    doc = Document()
    style_doc_title(doc, "CEO Meeting Agenda")
    style_doc_subtitle(doc, "Medicronis — Distributor Sales Reporting Platform · Phase 1 Approval")

    add_body(doc, "Duration: 45 minutes")
    add_body(doc, "Attendees: CEO, [project sponsor], [daily user / ops lead], MITRIXS (Abid Ali)")
    doc.add_paragraph()

    add_table_doc(
        doc,
        ["Time", "Topic"],
        [
            ["0–5 min", "Problem recap — manual PDF processing, daily FTE cost"],
            ["5–15 min", "Trial results & live recap (5 min demo optional)"],
            ["15–25 min", "Phase 1 solution, scope, rollout plan"],
            ["25–35 min", "Business case, investment, success metrics"],
            ["35–45 min", "Q&A, decision, next steps"],
        ],
    )

    add_section_heading(doc, "Materials Attached")
    for item in ["Executive one-pager", "Phase 1 scope", "Commercial summary"]:
        add_body(doc, item, bullet=True)

    add_section_heading(doc, "Decision Needed")
    add_body(doc, "Approve Phase 1 development and milestone-based engagement")

    path = ensure_output_dir() / "Meeting-Agenda.docx"
    doc.save(str(path))
    return path


def create_qa_cheat_sheet() -> Path:
    doc = Document()
    style_doc_title(doc, "CEO Q&A Cheat Sheet")
    style_doc_subtitle(doc, "Internal reference for MITRIXS — not for distribution")

    qa_pairs = [
        (
            "Why not fix our current sales software?",
            "It doesn't support distributor PDF ingestion. We build the missing layer; integration can follow.",
        ),
        (
            "Can you handle all 48 distributors?",
            "Yes, incrementally. Go-live with highest-volume templates; add 2–4 per week. Trial proved the engine; templates are configuration work.",
        ),
        (
            "What if accuracy isn't 100%?",
            "Neither is manual entry. Exceptions go to a review screen; approved mappings are reused. Target ≥85–90% auto-match on mature templates.",
        ),
        (
            "What about data security?",
            "Role-based access, encrypted storage, audit logs, HTTPS. Hosting per your policy. Phase 1 includes enterprise baseline security.",
        ),
        (
            "Why MITRIXS?",
            "We already built a working demo on your PDFs and SSR. Your team tested it. Phase 1 is a natural extension, not a greenfield gamble.",
        ),
        (
            "What if we only want reporting, not ERP?",
            "Phase 1 is reporting only. ERP modules are optional future phases — no commitment today.",
        ),
        (
            "What's the risk?",
            "Low: POC done, phased payments, clear acceptance criteria, UAT before final payment.",
        ),
    ]

    for question, answer in qa_pairs:
        add_section_heading(doc, f"Q: {question}")
        add_body(doc, f"A: {answer}")

    path = ensure_output_dir() / "CEO-QA-Cheat-Sheet.docx"
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# PDF helpers (ReportLab)
# ---------------------------------------------------------------------------


def build_pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="DocTitle",
            parent=styles["Heading1"],
            fontSize=18,
            textColor=colors.HexColor("#1A568E"),
            alignment=TA_CENTER,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="DocSubtitle",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#64748B"),
            alignment=TA_CENTER,
            spaceAfter=16,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Section",
            parent=styles["Heading2"],
            fontSize=13,
            textColor=colors.HexColor("#1A568E"),
            spaceBefore=12,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BodyText2",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            alignment=TA_JUSTIFY,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="BulletItem",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            leftIndent=14,
            bulletIndent=6,
            spaceAfter=4,
        )
    )
    return styles


def docx_to_pdf_content(title: str, subtitle: str, sections: list[tuple[str, list[str]]], output_path: Path) -> None:
    styles = build_pdf_styles()
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    story = []
    story.append(Paragraph(title, styles["DocTitle"]))
    story.append(Paragraph(subtitle, styles["DocSubtitle"]))
    story.append(Spacer(1, 0.2 * cm))

    for heading, paragraphs in sections:
        story.append(Paragraph(heading, styles["Section"]))
        for text in paragraphs:
            if text.startswith("• "):
                story.append(Paragraph(text, styles["BulletItem"], bulletText="•"))
            else:
                story.append(Paragraph(text, styles["BodyText2"]))
        story.append(Spacer(1, 0.15 * cm))

    doc.build(story)


def create_pdfs() -> list[Path]:
    out = ensure_output_dir()
    paths = []

    docx_to_pdf_content(
        "MEDICRONIS — Automated Distributor Sales Reporting",
        f"Executive Summary · MITRIXS · {date.today().strftime('%d %B %Y')}",
        [
            (
                "Problem",
                [
                    "Medicronis receives sales reports from 48+ distributors as PDFs in inconsistent formats. "
                    "Staff manually process these documents full-time before SSR and management reports can be "
                    "produced — causing delays, errors, and no centralized sales view.",
                ],
            ),
            (
                "Solution (Phase 1)",
                [
                    "MITRIXS proposes a Distributor Sales Intelligence Platform that ingests PDFs, extracts and "
                    "standardizes data, maps products to your master catalog, and generates SSR automatically — "
                    "with human review only for exceptions.",
                ],
            ),
            (
                "Proof",
                [
                    f"A working demo was built on Medicronis sample PDFs and SSR format. Your team tested the system for 2–3 days ({TRIAL_DATES}).",
                    "• [Insert 2–3 bullet outcomes from trial]",
                    "• [Insert user feedback quote]",
                ],
            ),
            (
                "Phase 1 Deliverables (6–8 weeks)",
                [
                    "PDF upload & processing · Product matching · Validation workflow · SSR generation · "
                    "Excel/PDF export · Basic dashboard · User roles & audit trail · Go-live with top 10–15 distributor templates",
                ],
            ),
            (
                "Investment",
                ["[Total amount] in three milestones: 30% kickoff / 40% core delivery / 30% go-live"],
            ),
            (
                "Expected Return",
                ["Recover ~150+ staff hours/month · Same-day SSR · Reduced mapping errors · Foundation for future sales analytics"],
            ),
            (
                "Decision Requested",
                ["Approve Phase 1 to begin formal development and production deployment."],
            ),
        ],
        out / "Executive-One-Pager.pdf",
    )
    paths.append(out / "Executive-One-Pager.pdf")

    docx_to_pdf_content(
        "Phase 1 — Scope of Work",
        "Medicronis Distributor Sales Intelligence Platform · Version 1.0",
        [
            ("1. Background", ["Medicronis requires automation of distributor PDF sales report processing and SSR generation. Phase 1 addresses this specific operational gap without replacing existing systems."]),
            ("2. Objectives", [
                "• Eliminate manual re-entry of distributor PDF data where possible",
                "• Generate SSR in Medicronis-approved format",
                "• Provide exception-based human validation",
                "• Establish centralized sales data for reporting",
            ]),
            ("3. Functional Requirements", [
                "User Management: Login, roles (Admin, Operator, Viewer), activity logging.",
                "Master Data: Distributor profiles, document templates, product master import.",
                "Document Processing: PDF upload, OCR, distributor identification.",
                "Product Matching: Exact, fuzzy, and review queue with saved mappings.",
                "Reporting: SSR per agreed template, Excel/PDF export.",
                "Dashboard: KPIs, processing status, exceptions pending.",
            ]),
            ("4. Client Responsibilities", [
                "• Provide SSR template, product master, representative PDFs",
                "• Assign UAT contact and provide timely feedback",
            ]),
            ("5. Out of Scope", [
                "• Inventory, procurement, finance, HR modules",
                "• Replacement of existing sales/accounting software",
                "• All 48 distributor templates at day-one go-live",
            ]),
            ("6. Timeline", ["Week 1: Kickoff · Weeks 2–6: Development · Weeks 7–8: UAT · End Week 8: Go-live"]),
            ("7. Acceptance Criteria", [
                "• SSR output matches agreed format",
                "• Go-live distributor template set processed successfully",
                "• User training and documentation delivered",
            ]),
        ],
        out / "Phase-1-Scope-of-Work.pdf",
    )
    paths.append(out / "Phase-1-Scope-of-Work.pdf")

    docx_to_pdf_content(
        "Medicronis Trial Results",
        TRIAL_DATES,
        [
            ("Participants", ["[Names / roles]"]),
            ("PDFs Tested", ["Complete the table in the Word version with distributor names, format types, and results."]),
            ("SSR Comparison", ["Compare manual SSR vs system SSR: row count, totals, sample products."]),
            ("Time Comparison", ["Process 1 PDF: [X] min manual vs [Y] min system. Generate SSR batch: [X] hrs vs [Y] min."]),
            ("Product Matching", ["Auto-matched: [X]% · Reviewed: [X]% · Unknown: [X] products"]),
            ("User Feedback", ["What worked · What needs improvement · Daily use willingness · Quote for CEO slide"]),
            ("Fix Plan for Phase 1", ["[How each issue will be addressed in MVP]"]),
        ],
        out / "Trial-Results-Template.pdf",
    )
    paths.append(out / "Trial-Results-Template.pdf")

    docx_to_pdf_content(
        "CEO Meeting Agenda",
        "Medicronis — Distributor Sales Reporting Platform · Phase 1 Approval",
        [
            ("Meeting Details", ["Duration: 45 minutes", "Attendees: CEO, project sponsor, daily user, MITRIXS (Abid Ali)"]),
            ("Agenda", [
                "0–5 min: Problem recap",
                "5–15 min: Trial results & demo recap",
                "15–25 min: Phase 1 solution, scope, rollout",
                "25–35 min: Business case, investment, KPIs",
                "35–45 min: Q&A, decision, next steps",
            ]),
            ("Decision Needed", ["Approve Phase 1 development and milestone-based engagement"]),
        ],
        out / "Meeting-Agenda.pdf",
    )
    paths.append(out / "Meeting-Agenda.pdf")

    docx_to_pdf_content(
        "CEO Q&A Cheat Sheet",
        "Internal reference for MITRIXS — not for distribution",
        [
            ("Why not fix current sales software?", ["It doesn't support distributor PDF ingestion. We build the missing layer; integration can follow."]),
            ("Can you handle all 48 distributors?", ["Yes, incrementally. Go-live with highest-volume templates; add 2–4 per week."]),
            ("What if accuracy isn't 100%?", ["Exceptions go to review; mappings are reused. Target ≥85–90% auto-match on mature templates."]),
            ("What about data security?", ["RBAC, encrypted storage, audit logs, HTTPS. Enterprise baseline in Phase 1."]),
            ("Why MITRIXS?", ["Working demo on your PDFs and SSR; your team tested it. Low-risk Phase 1 extension."]),
            ("ERP concerns?", ["Phase 1 is reporting only. Future modules are optional."]),
            ("What's the risk?", ["Low: POC done, phased payments, UAT before final payment."]),
        ],
        out / "CEO-QA-Cheat-Sheet.pdf",
    )
    paths.append(out / "CEO-QA-Cheat-Sheet.pdf")

    return paths


def create_presentation_pdf() -> Path:
    """Export slide deck content as a PDF (for sharing when PPTX isn't ideal)."""
    from reportlab.lib.pagesizes import landscape
    from reportlab.pdfgen import canvas

    out = ensure_output_dir() / "CEO-Presentation.pdf"
    page_w, page_h = landscape(A4)
    c = canvas.Canvas(str(out), pagesize=landscape(A4))

    slides = [
        (
            "Automated Distributor Sales Reporting",
            [
                "From PDF chaos to accurate SSR in minutes",
                "",
                "Prepared for: Medicronis",
                "Prepared by: MITRIXS — Abid Ali",
                f"Date: {MEETING_DATE}",
            ],
            True,
        ),
        (
            "The Problem Today",
            [
                "• 48+ distributors send sales data as PDF files in inconsistent formats",
                "• One employee spends all day, every day on manual PDF processing",
                "• SSR and management reports delayed until manual entry is complete",
                "• High error risk and no centralized sales view",
            ],
            False,
        ),
        (
            "What We Demonstrated (Trial Proof)",
            [
                "PDF processing: Manual all day → [X] minutes per batch",
                "SSR generation: After manual entry → Automated after upload",
                "Product matching: Manual → [X]% auto-matched",
                "Formats tested: [X] of [Y] PDFs successful",
                "",
                "User quote: \"[Insert feedback from daily user]\"",
            ],
            False,
        ),
        (
            "Proposed Solution — Phase 1",
            [
                "Distributor PDFs → Extract → Match products → SSR & reports",
                "Human review only for exceptions",
                "",
                "Upload PDFs · OCR for scanned docs · Product master mapping",
                "SSR in existing format · Excel/PDF export · Basic dashboard",
                "Works alongside current sales software",
            ],
            False,
        ),
        (
            "Phase 1 Scope",
            [
                "In scope (6–8 weeks): Auth, distributor/product master, PDF processing,",
                "matching + review, SSR, exports, dashboard, audit trail",
                "",
                "Out of scope: Inventory, procurement, finance, HR, ERP replacement",
            ],
            False,
        ),
        (
            "Rollout Plan",
            [
                "✓ Discovery + POC + Trial demo — Done",
                "Week 1–2: Requirements, product master, top 10 templates",
                "Week 3–6: Core platform and SSR workflow",
                "Week 7–8: UAT and go-live",
                "Post go-live: 2–4 distributor templates per week",
            ],
            False,
        ),
        (
            "Business Case / ROI",
            [
                "Current: ~176 staff hours/month on PDF/SSR processing",
                "Target: <20 hrs/month (exception review only)",
                "Time recovered: ~150+ hrs/month",
                "Same-day SSR vs multi-day delay today",
            ],
            False,
        ),
        (
            "Investment & Commercial Terms",
            [
                "Kickoff (30%): Signed scope, access, samples — [Amount]",
                "Core engine (40%): Upload, extract, match, review — [Amount]",
                "Go-live (30%): SSR production, [X] templates — [Amount]",
                "",
                "Total Phase 1: [Total Amount]",
            ],
            False,
        ),
        (
            "Success Metrics",
            [
                "≥90% extraction accuracy on digital PDFs",
                "≥85% product auto-match rate",
                "≥80% reduction in manual data entry hours",
                "10–15 distributor templates at go-live",
                "Review at 30 / 60 / 90 days",
            ],
            False,
        ),
        (
            "Recommended Decision & Next Steps",
            [
                "1. Approve Phase 1 scope & commercial terms",
                "2. Nominate project sponsor + UAT contact",
                "3. Provide product master + remaining PDF samples",
                "4. Kickoff within [X] days of signing",
                "",
                "Abid Ali · MITRIXS · 03204014920",
            ],
            False,
        ),
    ]

    for title, lines, is_title_slide in slides:
        if is_title_slide:
            c.setFillColor(colors.HexColor("#1A568E"))
            c.rect(0, 0, page_w, page_h, fill=1, stroke=0)
            c.setFillColor(colors.white)
            c.setFont("Helvetica-Bold", 28)
            c.drawString(2 * cm, page_h - 4 * cm, title)
            c.setFont("Helvetica", 14)
            y = page_h - 5.5 * cm
            for line in lines:
                c.drawString(2 * cm, y, line)
                y -= 0.7 * cm
        else:
            c.setFillColor(colors.white)
            c.rect(0, 0, page_w, page_h, fill=1, stroke=0)
            c.setFillColor(colors.HexColor("#1A568E"))
            c.rect(0, page_h - 2 * cm, page_w, 2 * cm, fill=1, stroke=0)
            c.setFillColor(colors.white)
            c.setFont("Helvetica-Bold", 18)
            c.drawString(1.5 * cm, page_h - 1.45 * cm, title)
            c.setFillColor(colors.HexColor("#1E293B"))
            c.setFont("Helvetica", 12)
            y = page_h - 3 * cm
            for line in lines:
                c.drawString(1.8 * cm, y, line)
                y -= 0.65 * cm
        c.showPage()

    c.save()
    return out


def main() -> None:
    ensure_output_dir()
    created = []

    ppt = create_presentation()
    created.append(ppt)
    created.append(create_presentation_pdf())

    for creator in [
        create_executive_one_pager,
        create_phase1_scope,
        create_trial_results_template,
        create_meeting_agenda,
        create_qa_cheat_sheet,
    ]:
        created.append(creator())

    created.extend(create_pdfs())

    print("Generated files:")
    for path in created:
        print(f"  {path}")


if __name__ == "__main__":
    main()
